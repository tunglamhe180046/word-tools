"""
Tests for tools/word-engine/cli.py - the --json stdout/exit-code contract described in
docs/plans/DEDICATED_WORD_ENGINE_TOOL_PLAN.md muc 1.2 and 2.

Real subprocess.run() calls against cli.py (not in-process imports): this is exactly how the
TypeScript consumers (phan-tich/, "nhan vien ho so/") and Claude/Jarvis are meant to invoke this
tool (child_process.execFile("python", ["tools/word-engine/cli.py", "--json", ...])), so the test
should exercise the real process boundary - argv parsing, stdout/stderr separation, exit code -
not just the Python functions cli.py wraps (those already have their own tests in
tests/test_surgical_safety.py and tests/test_safe_restore.py). Every invocation passes an explicit
--work-dir tmp_path so `.jarvis/` never touches the real repo's `.jarvis/`.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document

CLI_PATH = Path(__file__).resolve().parent.parent / "cli.py"


# ---------------------------------------------------------------------------
# Fixtures & small helpers
# ---------------------------------------------------------------------------
def _build_sample_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Doan van thu nhat.")
    doc.add_paragraph("Doan van thu hai.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "O 1"
    table.cell(0, 1).text = "8.5"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _run_cli(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(CLI_PATH), *args],
        capture_output=True,
        text=True,
    )


def _inspect_json(docx_path: Path, tmp_path: Path) -> dict:
    proc = _run_cli("inspect", str(docx_path), "--json", "--work-dir", str(tmp_path))
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout)


# ---------------------------------------------------------------------------
# inspect
# ---------------------------------------------------------------------------
def test_cli_inspect_json_contract(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    proc = _run_cli("inspect", str(docx_path), "--json", "--work-dir", str(tmp_path))

    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.count("\n") == 1  # duy nhat 1 dong JSON, khong co log rac khac
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["document_revision"].startswith("sha256:")
    object_ids = {loc["object_id"] for loc in payload["locators"]}
    assert "para_1" in object_ids
    assert "cell_t1_r1_c2" in object_ids


# ---------------------------------------------------------------------------
# patch-cell
# ---------------------------------------------------------------------------
def test_cli_patch_cell_updates_docx_and_returns_json(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    revision = _inspect_json(docx_path, tmp_path)["document_revision"]

    proc = _run_cli(
        "patch-cell",
        str(docx_path),
        "--target-id",
        "cell_t1_r1_c2",
        "--new-text",
        "9.0",
        "--expected-revision",
        revision,
        "--json",
        "--work-dir",
        str(tmp_path),
        "--job-id",
        "job-cell",
    )

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["outcome"] == "committed_clean"
    assert payload["document_revision"].startswith("sha256:")
    assert payload["document_revision"] != revision

    doc = Document(str(docx_path))
    assert doc.tables[0].cell(0, 1).text == "9.0"


# ---------------------------------------------------------------------------
# patch-text
# ---------------------------------------------------------------------------
def test_cli_patch_text_without_locator_updates_docx(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    proc = _run_cli(
        "patch-text",
        str(docx_path),
        "--search",
        "Doan van thu nhat.",
        "--replace",
        "Doan van da duoc sua.",
        "--json",
        "--work-dir",
        str(tmp_path),
    )

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["outcome"] == "committed_clean"

    doc = Document(str(docx_path))
    assert doc.paragraphs[0].text == "Doan van da duoc sua."


def test_cli_patch_text_with_locator_scopes_to_target_cell(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    revision = _inspect_json(docx_path, tmp_path)["document_revision"]

    proc = _run_cli(
        "patch-text",
        str(docx_path),
        "--search",
        "O 1",
        "--replace",
        "O 1 (da sua)",
        "--target-id",
        "cell_t1_r1_c1",
        "--expected-revision",
        revision,
        "--json",
        "--work-dir",
        str(tmp_path),
    )

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True

    doc = Document(str(docx_path))
    assert doc.tables[0].cell(0, 0).text == "O 1 (da sua)"


# ---------------------------------------------------------------------------
# backups
# ---------------------------------------------------------------------------
def test_cli_backups_lists_backup_created_by_patch(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    revision = _inspect_json(docx_path, tmp_path)["document_revision"]

    patch_proc = _run_cli(
        "patch-cell",
        str(docx_path),
        "--target-id",
        "cell_t1_r1_c2",
        "--new-text",
        "9.0",
        "--expected-revision",
        revision,
        "--json",
        "--work-dir",
        str(tmp_path),
    )
    assert patch_proc.returncode == 0, patch_proc.stderr

    proc = _run_cli("backups", str(docx_path), "--json", "--work-dir", str(tmp_path))

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert len(payload["backups"]) == 1
    assert payload["backups"][0]["sha256"]


def test_cli_backups_empty_before_any_write(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    proc = _run_cli("backups", str(docx_path), "--json", "--work-dir", str(tmp_path))

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["backups"] == []


# ---------------------------------------------------------------------------
# restore
# ---------------------------------------------------------------------------
def test_cli_restore_reverts_content(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    revision = _inspect_json(docx_path, tmp_path)["document_revision"]

    patch_proc = _run_cli(
        "patch-cell",
        str(docx_path),
        "--target-id",
        "cell_t1_r1_c2",
        "--new-text",
        "9.0",
        "--expected-revision",
        revision,
        "--json",
        "--work-dir",
        str(tmp_path),
    )
    assert patch_proc.returncode == 0, patch_proc.stderr

    backups_payload = json.loads(_run_cli("backups", str(docx_path), "--json", "--work-dir", str(tmp_path)).stdout)
    backup_id = backups_payload["backups"][0]["backup_id"]

    proc = _run_cli(
        "restore",
        str(docx_path),
        "--backup-id",
        backup_id,
        "--json",
        "--work-dir",
        str(tmp_path),
    )

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["outcome"] == "restored_from_backup"

    doc = Document(str(docx_path))
    assert doc.tables[0].cell(0, 1).text == "8.5"


def test_cli_restore_unknown_backup_id_fails_cleanly(tmp_path):
    # restore_backup() only accepts a target that is already tracked by the broker (has gone
    # through issue_job/request_commit at least once) - an untracked file fails with
    # RestoreConflictError before backup_id is even looked up (see tests/test_safe_restore.py's
    # test_restore_backup_raises_conflict_when_target_never_tracked). So this test must first put
    # the docx through one patch (tracks it, creates baseline) before asking for an unknown
    # backup_id, to exercise BackupNotFoundError specifically.
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    revision = _inspect_json(docx_path, tmp_path)["document_revision"]
    patch_proc = _run_cli(
        "patch-cell",
        str(docx_path),
        "--target-id",
        "cell_t1_r1_c2",
        "--new-text",
        "9.0",
        "--expected-revision",
        revision,
        "--json",
        "--work-dir",
        str(tmp_path),
    )
    assert patch_proc.returncode == 0, patch_proc.stderr

    proc = _run_cli(
        "restore",
        str(docx_path),
        "--backup-id",
        "does-not-exist",
        "--json",
        "--work-dir",
        str(tmp_path),
    )

    assert proc.returncode == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "BackupNotFoundError"


# ---------------------------------------------------------------------------
# Error contract: exit code 1 + parseable {"success": false, ...} JSON on stdout
# ---------------------------------------------------------------------------
def test_cli_patch_cell_wrong_expected_revision_fails_cleanly_and_leaves_disk_untouched(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    before_bytes = docx_path.read_bytes()

    proc = _run_cli(
        "patch-cell",
        str(docx_path),
        "--target-id",
        "cell_t1_r1_c2",
        "--new-text",
        "9.0",
        "--expected-revision",
        "sha256:0000000000000000000000000000000000000000000000000000000000000000",
        "--json",
        "--work-dir",
        str(tmp_path),
    )

    assert proc.returncode == 1
    assert proc.stdout.count("\n") == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "DocumentDriftError"
    assert payload["reason"]
    assert docx_path.read_bytes() == before_bytes


def test_cli_patch_text_search_not_found_fails_cleanly_and_leaves_disk_untouched(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    before_bytes = docx_path.read_bytes()

    proc = _run_cli(
        "patch-text",
        str(docx_path),
        "--search",
        "Khong ton tai trong tai lieu nay",
        "--replace",
        "X",
        "--json",
        "--work-dir",
        str(tmp_path),
    )

    assert proc.returncode == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "TextNotFoundError"
    assert docx_path.read_bytes() == before_bytes
