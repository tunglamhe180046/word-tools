"""
Tests for core/surgical_patcher.py - OOXML Surgical Patcher (Profile 1), see
docs/plans/DEDICATED_WORD_ENGINE_TOOL_PLAN.md muc 3.3 and 5.1 for the measurable invariants these
tests verify, and core/surgical_patcher.py's module docstring for the design.

No mocking, real files on disk under tmp_path, real SHA-256/C14N comparisons - same style as
tests/test_locators.py and tests/test_safe_restore.py. Every test passes an explicit
work_dir=tmp_path so `.jarvis/` never touches the real repo's `.jarvis/`.
"""
from __future__ import annotations

import zipfile
from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from core.inspector import (
    WORD_NS,
    DocumentDriftError,
    ObjectLocator,
    inspect_document,
    load_document_root,
)
from core.safety_gateway import sha256_file
from core.surgical_patcher import (
    AmbiguousTextMatchError,
    CellSpansMultipleParagraphsError,
    patch_cell,
    patch_text,
)
from core.text_resolver import TextNotFoundError

_NSMAP = {"w": WORD_NS}
_T_TAG = f"{{{WORD_NS}}}t"


# ---------------------------------------------------------------------------
# Fixtures & small helpers
# ---------------------------------------------------------------------------
def _build_sample_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Doan van thu nhat.")
    doc.add_paragraph("Doan van thu hai.")
    doc.add_paragraph("Doan van thu ba.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "O 1"
    table.cell(0, 1).text = "O 2"
    table.cell(1, 0).text = "O 3"
    table.cell(1, 1).text = "8.5"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _locator_dict(report: dict, object_id: str) -> dict:
    return next(loc for loc in report["locators"] if loc["object_id"] == object_id)


def _locator(report: dict, object_id: str) -> ObjectLocator:
    return ObjectLocator(**_locator_dict(report, object_id))


def _cell_text(body: etree._Element, table_index: int, row_index: int, col_index: int) -> str:
    tbl = body.findall("w:tbl", _NSMAP)[table_index]
    tc = tbl.findall("w:tr", _NSMAP)[row_index].findall("w:tc", _NSMAP)[col_index]
    return "".join(t.text or "" for t in tc.iter(_T_TAG))


def _part_hashes(docx_path: Path) -> dict:
    with zipfile.ZipFile(docx_path, "r") as zf:
        return {name: sha256(zf.read(name)).hexdigest() for name in zf.namelist()}


COMMON_KWARGS = dict(actor="tester", job_id="job1")


# ---------------------------------------------------------------------------
# Profile 1 invariant: Untouched OPC parts RAW SHA-256 IDENTICAL 100%
# ---------------------------------------------------------------------------
def test_patch_cell_leaves_untouched_opc_parts_raw_sha256_identical(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r2_c2")

    before = _part_hashes(docx_path)

    result = patch_cell(docx_path, locator, "9.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)
    assert result.success is True

    after = _part_hashes(docx_path)
    assert set(before) == set(after)
    for name in before:
        if name == "word/document.xml":
            continue
        assert after[name] == before[name], f"part {name} changed unexpectedly"
    assert after["word/document.xml"] != before["word/document.xml"]


# ---------------------------------------------------------------------------
# Profile 1 invariant: word/document.xml untouched subtrees C14N-identical,
# only the target subtree has a semantic diff
# ---------------------------------------------------------------------------
def test_patch_cell_leaves_other_subtrees_c14n_identical_and_only_target_diffs(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r2_c2")  # row_index=1, col_index=1

    root_before = load_document_root(docx_path)
    body_before = root_before.find("w:body", _NSMAP)
    paras_before = body_before.findall("w:p", _NSMAP)
    tbls_before = body_before.findall("w:tbl", _NSMAP)

    patch_cell(docx_path, locator, "9.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)

    root_after = load_document_root(docx_path)
    body_after = root_after.find("w:body", _NSMAP)
    paras_after = body_after.findall("w:p", _NSMAP)
    tbls_after = body_after.findall("w:tbl", _NSMAP)

    # Every top-level paragraph is untouched.
    assert len(paras_before) == len(paras_after)
    for pb, pa in zip(paras_before, paras_after):
        assert etree.tostring(pb, method="c14n") == etree.tostring(pa, method="c14n")

    assert len(tbls_before) == len(tbls_after) == 1
    rows_before = tbls_before[0].findall("w:tr", _NSMAP)
    rows_after = tbls_after[0].findall("w:tr", _NSMAP)
    for ri, (rb, ra) in enumerate(zip(rows_before, rows_after)):
        cells_before = rb.findall("w:tc", _NSMAP)
        cells_after = ra.findall("w:tc", _NSMAP)
        for ci, (cb, ca) in enumerate(zip(cells_before, cells_after)):
            is_target = ri == 1 and ci == 1
            c14n_before = etree.tostring(cb, method="c14n")
            c14n_after = etree.tostring(ca, method="c14n")
            if is_target:
                assert c14n_before != c14n_after
                assert "".join(t.text or "" for t in ca.iter(_T_TAG)) == "9.0"
            else:
                assert c14n_before == c14n_after, f"cell (row={ri}, col={ci}) changed unexpectedly"


def test_patch_cell_preserves_w_tcPr(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    # Give the target cell a distinguishing tcPr so "preserved" is a meaningful assertion, not a
    # trivial "empty matches empty".
    doc = Document(str(docx_path))
    tc = doc.tables[0].cell(1, 1)._tc
    tcPr = tc.find(f"{{{WORD_NS}}}tcPr")
    if tcPr is None:
        tcPr = etree.SubElement(tc, f"{{{WORD_NS}}}tcPr")
        tc.insert(0, tcPr)
    shd = etree.SubElement(tcPr, f"{{{WORD_NS}}}shd")
    shd.set(f"{{{WORD_NS}}}fill", "FFFF00")
    doc.save(docx_path)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r2_c2")

    root_before = load_document_root(docx_path)
    body_before = root_before.find("w:body", _NSMAP)
    tc_before = body_before.findall("w:tbl", _NSMAP)[0].findall("w:tr", _NSMAP)[1].findall("w:tc", _NSMAP)[1]
    tcPr_before = etree.tostring(tc_before.find("w:tcPr", _NSMAP), method="c14n")

    patch_cell(docx_path, locator, "9.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)

    root_after = load_document_root(docx_path)
    body_after = root_after.find("w:body", _NSMAP)
    tc_after = body_after.findall("w:tbl", _NSMAP)[0].findall("w:tr", _NSMAP)[1].findall("w:tc", _NSMAP)[1]
    tcPr_after = etree.tostring(tc_after.find("w:tcPr", _NSMAP), method="c14n")

    assert tcPr_before == tcPr_after
    assert "".join(t.text or "" for t in tc_after.iter(_T_TAG)) == "9.0"


# ---------------------------------------------------------------------------
# Safety Gateway integration: automatic backup before overwrite
# ---------------------------------------------------------------------------
def test_patch_cell_creates_backup_via_commit_broker_with_old_content(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r2_c2")

    result = patch_cell(docx_path, locator, "9.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)

    assert result.backup_path is not None
    backup_path = Path(result.backup_path)
    assert backup_path.exists()

    with zipfile.ZipFile(backup_path, "r") as zf:
        backup_root = etree.fromstring(zf.read("word/document.xml"))
    backup_body = backup_root.find("w:body", _NSMAP)
    assert _cell_text(backup_body, 0, 1, 1) == "8.5"


# ---------------------------------------------------------------------------
# State machine: DOCUMENT_DRIFT blocks the patch, disk left untouched
# ---------------------------------------------------------------------------
def test_patch_cell_raises_document_drift_when_disk_hand_edited(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r2_c2")

    doc = Document(str(docx_path))
    doc.add_paragraph("Nguoi dung vua sua tay them doan nay.")
    doc.save(docx_path)
    hand_edited_bytes = docx_path.read_bytes()

    with pytest.raises(DocumentDriftError):
        patch_cell(docx_path, locator, "9.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)

    assert docx_path.read_bytes() == hand_edited_bytes


def test_patch_text_raises_document_drift_when_disk_hand_edited(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "para_1")

    doc = Document(str(docx_path))
    doc.add_paragraph("Sua tay ngoai broker.")
    doc.save(docx_path)
    hand_edited_bytes = docx_path.read_bytes()

    with pytest.raises(DocumentDriftError):
        patch_text(
            docx_path, "Doan van thu nhat.", "Da sua", locator, None,
            work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
        )

    assert docx_path.read_bytes() == hand_edited_bytes


# ---------------------------------------------------------------------------
# Chaining: patch #1's new_locators let patch #2 land correctly, and the
# pre-patch1 (now-stale) locator is correctly rejected as drifted.
# ---------------------------------------------------------------------------
def test_chained_patches_use_returned_new_locators_without_drift(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator1 = _locator(report, "cell_t1_r2_c2")

    result1 = patch_cell(docx_path, locator1, "9.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)
    assert result1.success is True
    assert result1.outcome == "committed_clean"
    assert result1.new_document_revision == f"sha256:{sha256_file(docx_path)}"

    locator2_dict = next(loc for loc in result1.new_locators if loc["object_id"] == "para_1")
    locator2 = ObjectLocator(**locator2_dict)

    result2 = patch_text(
        docx_path, "Doan van thu nhat.", "Doan van da duoc sua.", locator2, None,
        work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )
    assert result2.success is True

    root_final = load_document_root(docx_path)
    body_final = root_final.find("w:body", _NSMAP)
    first_p_text = "".join(t.text or "" for t in body_final.findall("w:p", _NSMAP)[0].iter(_T_TAG))
    assert first_p_text == "Doan van da duoc sua."
    # Patch #1's edit survived patch #2 untouched (chaining did not clobber earlier work).
    assert _cell_text(body_final, 0, 1, 1) == "9.0"

    # The pre-patch1 locator's revision is now stale - reusing it must be rejected, never
    # silently "rebased" (plan muc 3.1: loai bo hoan toan co che tu rebase locator).
    with pytest.raises(DocumentDriftError):
        patch_cell(docx_path, locator1, "0.0", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)


# ---------------------------------------------------------------------------
# patch_cell(): scope guardrails (Phase 1 MVP limits)
# ---------------------------------------------------------------------------
def test_patch_cell_rejects_paragraph_locator(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "para_1")

    with pytest.raises(ValueError):
        patch_cell(docx_path, locator, "X", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)


def test_patch_cell_raises_when_cell_has_multiple_paragraphs(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "Dong 1"
    cell.add_paragraph("Dong 2")
    doc.save(docx_path)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r1_c1")
    before_bytes = docx_path.read_bytes()

    with pytest.raises(CellSpansMultipleParagraphsError):
        patch_cell(docx_path, locator, "X", work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS)

    assert docx_path.read_bytes() == before_bytes


# ---------------------------------------------------------------------------
# patch_text(): locator-scoped table_cell search, and document-wide search
# ---------------------------------------------------------------------------
def test_patch_text_with_table_cell_locator_replaces_substring(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = _locator(report, "cell_t1_r1_c1")  # "O 1"

    result = patch_text(
        docx_path, "O 1", "O 1 (da sua)", locator, None,
        work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )
    assert result.success is True
    root = load_document_root(docx_path)
    body = root.find("w:body", _NSMAP)
    assert _cell_text(body, 0, 0, 0) == "O 1 (da sua)"


def test_patch_text_without_locator_finds_unique_paragraph(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    result = patch_text(
        docx_path, "Doan van thu hai.", "Doan van thu hai da sua.", None, None,
        work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )
    assert result.success is True

    root = load_document_root(docx_path)
    body = root.find("w:body", _NSMAP)
    texts = ["".join(t.text or "" for t in p.iter(_T_TAG)) for p in body.findall("w:p", _NSMAP)]
    assert "Doan van thu hai da sua." in texts


def test_patch_text_without_locator_raises_ambiguous_when_multiple_paragraphs_match(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Trung lap noi dung nay.")
    doc.add_paragraph("Trung lap noi dung nay.")
    doc.save(docx_path)
    before_bytes = docx_path.read_bytes()

    with pytest.raises(AmbiguousTextMatchError):
        patch_text(
            docx_path, "Trung lap", "Da sua", None, None,
            work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
        )

    assert docx_path.read_bytes() == before_bytes


def test_patch_text_raises_text_not_found_when_absent(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    before_bytes = docx_path.read_bytes()

    with pytest.raises(TextNotFoundError):
        patch_text(
            docx_path, "Khong ton tai trong tai lieu nay", "X", None, None,
            work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
        )

    assert docx_path.read_bytes() == before_bytes
