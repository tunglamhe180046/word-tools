"""
Tests for core/inspector.py - Revision-bound Object Locators (see
docs/plans/DEDICATED_WORD_ENGINE_TOOL_PLAN.md section 3.1 and core/inspector.py's module
docstring for the design, especially which fields are deliberately excluded from
context_sha256's payload so fallback re-scan can survive a positional shift).

Fixtures are built with python-docx (already an independent Phase 1 dependency, see
requirements.txt) rather than hand-written XML strings, per the task's own instruction. No
mocking, real files on disk under tmp_path - same style as tests/test_safe_restore.py.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from core.inspector import (
    WORD_NS,
    AmbiguousLocatorError,
    DocumentDriftError,
    LocatorNotFoundError,
    inspect_document,
    load_document_root,
    resolve_in_document,
    resolve_locator,
)
from core.safety_gateway import sha256_file

_SHA256_HEX_RE = re.compile(r"^[0-9a-f]{64}$")


def _build_sample_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Doan van thu nhat.")
    doc.add_paragraph("Doan van thu hai.")
    doc.add_paragraph("Doan van thu ba.")
    doc.add_paragraph("Doan van thu tu.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "O 1"
    table.cell(0, 1).text = "O 2"
    table.cell(1, 0).text = "O 3"
    table.cell(1, 1).text = "O 4"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _make_paragraph_element(text: str) -> etree._Element:
    p = etree.Element(f"{{{WORD_NS}}}p")
    r = etree.SubElement(p, f"{{{WORD_NS}}}r")
    t = etree.SubElement(r, f"{{{WORD_NS}}}t")
    t.text = text
    return p


def _make_row_element(texts: list[str]) -> etree._Element:
    tr = etree.Element(f"{{{WORD_NS}}}tr")
    for text in texts:
        tc = etree.SubElement(tr, f"{{{WORD_NS}}}tc")
        p = etree.SubElement(tc, f"{{{WORD_NS}}}p")
        r = etree.SubElement(p, f"{{{WORD_NS}}}r")
        t = etree.SubElement(r, f"{{{WORD_NS}}}t")
        t.text = text
    return tr


# ---------------------------------------------------------------------------
# inspect_document(): document_revision, locators, structural_path
# ---------------------------------------------------------------------------
def test_inspect_document_generates_document_revision_and_locators(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")

    assert report["document_revision"] == f"sha256:{sha256_file(docx_path)}"

    paragraph_locators = [loc for loc in report["locators"] if loc["kind"] == "paragraph"]
    assert [loc["structural_path"] for loc in paragraph_locators] == [
        "/w:document/w:body/w:p[1]",
        "/w:document/w:body/w:p[2]",
        "/w:document/w:body/w:p[3]",
        "/w:document/w:body/w:p[4]",
    ]
    assert [loc["expected_text"] for loc in paragraph_locators] == [
        "Doan van thu nhat.",
        "Doan van thu hai.",
        "Doan van thu ba.",
        "Doan van thu tu.",
    ]
    assert [loc["object_id"] for loc in paragraph_locators] == ["para_1", "para_2", "para_3", "para_4"]

    cell_locators = [loc for loc in report["locators"] if loc["kind"] == "table_cell"]
    assert len(cell_locators) == 4
    last_cell = next(loc for loc in cell_locators if loc["expected_text"] == "O 4")
    assert last_cell["structural_path"] == "/w:document/w:body/w:tbl[1]/w:tr[2]/w:tc[2]"
    assert last_cell["object_id"] == "cell_t1_r2_c2"
    assert last_cell["table_index"] == 0
    assert last_cell["row_index"] == 1
    assert last_cell["col_index"] == 1


def test_context_sha256_is_valid_sha256_hex_for_every_locator(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")

    assert report["locators"], "fixture must produce at least one locator"
    for locator in report["locators"]:
        assert locator["context_sha256"].startswith("sha256:")
        assert _SHA256_HEX_RE.match(locator["context_sha256"].removeprefix("sha256:"))


# ---------------------------------------------------------------------------
# Sidecar index file
# ---------------------------------------------------------------------------
def test_sidecar_index_written_with_expected_structure(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job-abc")

    index_path = tmp_path / ".jarvis" / "work" / "job-abc" / "document-index.json"
    assert index_path.exists()
    assert report["index_file_path"] == str(index_path)

    data = json.loads(index_path.read_text(encoding="utf-8"))
    assert data["schema"] == "word-engine-document-index.v1"
    assert data["job_id"] == "job-abc"
    assert data["docx_path"] == str(docx_path.resolve())
    assert data["document_revision"] == report["document_revision"]
    assert data["locators"] == report["locators"]


# ---------------------------------------------------------------------------
# Fast structural lookup
# ---------------------------------------------------------------------------
def test_fast_structural_lookup_finds_correct_node(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")

    paragraph_locator = next(loc for loc in report["locators"] if loc["object_id"] == "para_3")
    element = resolve_locator(docx_path, paragraph_locator)
    assert "".join(t.text or "" for t in element.iter(f"{{{WORD_NS}}}t")) == "Doan van thu ba."

    cell_locator = next(loc for loc in report["locators"] if loc["object_id"] == "cell_t1_r2_c2")
    element = resolve_locator(docx_path, cell_locator)
    assert "".join(t.text or "" for t in element.iter(f"{{{WORD_NS}}}t")) == "O 4"


def test_resolve_locator_raises_document_drift_when_disk_changed(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    paragraph_locator = next(loc for loc in report["locators"] if loc["object_id"] == "para_1")

    # A hand-edit (or any other write) after inspect_document() changes the disk hash.
    doc = Document(str(docx_path))
    doc.add_paragraph("Doan van moi them vao.")
    doc.save(docx_path)

    with pytest.raises(DocumentDriftError):
        resolve_locator(docx_path, paragraph_locator)


# ---------------------------------------------------------------------------
# Fallback re-scan by context_sha256 when structural_path has drifted
# ---------------------------------------------------------------------------
def test_fallback_rescan_finds_paragraph_after_structural_path_drift(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    target_locator = next(loc for loc in report["locators"] if loc["object_id"] == "para_3")

    root = load_document_root(docx_path)
    body = root.find("w:body", {"w": WORD_NS})
    # Simulate drift from an internal transaction elsewhere in the document: insert a brand-new
    # paragraph before the very first one, shifting every later w:p[i] index by 1. The target's
    # own immediate neighbors (para_2, para_4) are untouched, so its context_sha256 still matches.
    body.insert(0, _make_paragraph_element("Doan van chen them."))

    # Fast lookup on the stale structural_path (w:p[3]) now lands on the wrong (shifted)
    # paragraph - the original second paragraph, not the target's text.
    shifted = body.findall("w:p", {"w": WORD_NS})[2]
    assert "".join(t.text or "" for t in shifted.iter(f"{{{WORD_NS}}}t")) == "Doan van thu hai."

    element = resolve_in_document(root, target_locator)
    assert "".join(t.text or "" for t in element.iter(f"{{{WORD_NS}}}t")) == "Doan van thu ba."


def test_fallback_rescan_finds_table_cell_after_row_insertion_drift(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    target_locator = next(loc for loc in report["locators"] if loc["object_id"] == "cell_t1_r2_c2")

    root = load_document_root(docx_path)
    body = root.find("w:body", {"w": WORD_NS})
    tbl = body.findall("w:tbl", {"w": WORD_NS})[0]
    # Insert a new row before the existing rows: the target cell's row_index shifts, but its
    # table_index and same-row neighbor text ("O 3") are unchanged, so context_sha256 still
    # matches per core/inspector.py's deliberate exclusion of row_index/col_index from the hash.
    tbl.insert(0, _make_row_element(["X1", "X2"]))

    element = resolve_in_document(root, target_locator)
    assert "".join(t.text or "" for t in element.iter(f"{{{WORD_NS}}}t")) == "O 4"


def test_fallback_rescan_raises_locator_not_found_when_content_removed(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    target_locator = next(loc for loc in report["locators"] if loc["object_id"] == "para_2")

    root = load_document_root(docx_path)
    body = root.find("w:body", {"w": WORD_NS})
    for p in body.findall("w:p", {"w": WORD_NS}):
        body.remove(p)

    with pytest.raises(LocatorNotFoundError):
        resolve_in_document(root, target_locator)


def test_fallback_rescan_raises_ambiguous_when_multiple_matches(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    target_locator = next(loc for loc in report["locators"] if loc["object_id"] == "para_2")

    root = load_document_root(docx_path)
    body = root.find("w:body", {"w": WORD_NS})

    # Insert an unrelated paragraph (shifts every later w:p[i] index, so the fast structural
    # lookup at the stale path misses) immediately followed by an exact duplicate of the
    # (prev, target, next) triplet. The rescan now finds two equally valid context_sha256
    # matches: this new duplicate triplet's middle paragraph, and the untouched original further
    # down - correctly refusing to guess which one is the real target.
    block_texts = [
        "Doan chen khong lien quan.",
        "Doan van thu nhat.",
        "Doan van thu hai.",
        "Doan van thu ba.",
    ]
    for offset, text in enumerate(block_texts):
        body.insert(offset, _make_paragraph_element(text))

    with pytest.raises(AmbiguousLocatorError):
        resolve_in_document(root, target_locator)
