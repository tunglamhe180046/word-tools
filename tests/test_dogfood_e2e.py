"""
tools/word-engine/tests/test_dogfood_e2e.py - End-to-End Dogfooding cho tools/word-engine/.

Khac voi cac test khac trong thu muc nay (moi file chi kiem thu DUNG 1 module), file nay dung
TAT CA cac module cong khai cua chinh du an (core/inspector.py, core/surgical_patcher.py,
core/geometry_styler.py, core/stamp_ops.py, core/safety_gateway.py) qua DUNG 1 kich ban phau
thuat hoan chinh tren MOT tai lieu .docx duy nhat, mo phong dung trinh tu mot nhan vien phong ho
so se lam tren mot bang diem that: sua 1 o diem, sua 1 cum tu trong doan ket luan, chuan hoa kho
giay/le/dan trang/vien bang, chen anh dau, kiem tra danh sach backup, roi phuc hoi - va xac nhan
tung buoc bang cach doc lai chinh file .docx that qua python-docx/lxml sau MOI thao tac (khong chi
tin vao JSON/PatchResult tra ve).

Goi truc tiep cac ham Python cong khai (khong qua subprocess CLI) - dung tinh than
tests/test_geometry.py va tests/test_stamp_ops.py: cli.py chi la lop dieu phoi mong ghep lai
chinh cac ham nay (xem cli.py module docstring), nen goi thang qua duong bien process that
(subprocess) da co rieng trong tests/test_cli.py.

No mocking, real files tren dia duoi tmp_path - work_dir=tmp_path tuong minh o MOI loi goi de
`.jarvis/` khong bao gio dung vao `.jarvis/` that cua repo (cung quy uoc voi moi test khac).
"""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any, Dict

import pytest
from docx import Document
from PIL import Image

from core.geometry_styler import apply_geometry
from core.inspector import WORD_NS, ObjectLocator, inspect_document, load_document_root
from core.safety_gateway import RestoreConflictError, list_backups, restore_backup
from core.stamp_ops import apply_stamp, register_stamp_asset
from core.surgical_patcher import patch_cell, patch_text

_NSMAP = {"w": WORD_NS}
ACTOR = "dogfood-tester"


# ---------------------------------------------------------------------------
# Fixtures & small helpers
# ---------------------------------------------------------------------------
def _build_report_card_docx(path: Path) -> None:
    """Tai lieu 'thuc te': tieu de + doan thong tin hoc sinh + bang diem 3 mon (Toan/Ly/Hoa) +
    doan van ket luan - dung cau truc de bai yeu cau ('Tieu de van ban', 'Bang bieu diem hoc tap',
    'Doan van ket luan')."""
    doc = Document()
    doc.add_heading("BANG DIEM HOC TAP", level=1)
    doc.add_paragraph("Ho va ten hoc sinh: Nguyen Van Dogfood - Lop 12A1.")
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Mon hoc"
    table.cell(0, 1).text = "Diem"
    table.cell(1, 0).text = "Toan"
    table.cell(1, 1).text = "8.5"
    table.cell(2, 0).text = "Ly"
    table.cell(2, 1).text = "7.5"
    table.cell(3, 0).text = "Hoa"
    table.cell(3, 1).text = "8.0"
    doc.add_paragraph(
        "Ket luan: Hoc sinh co ket qua hoc tap kha, du dieu kien xet tot nghiep."
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _make_stamp_png(path: Path) -> None:
    """Anh PNG that (khong phai byte gia) - dung Pillow, mot trong cac dependency da khai bao cua
    tools/word-engine/, de dung tinh than 'tai lieu thuc te' cua kich ban dogfood nay."""
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (60, 60), color=(200, 30, 30)).save(path, format="PNG")


def _locator(report: Dict[str, Any], object_id: str) -> ObjectLocator:
    return ObjectLocator(**next(loc for loc in report["locators"] if loc["object_id"] == object_id))


def _sect_pr(docx_path: Path):
    root = load_document_root(docx_path)
    body = root.find("w:body", _NSMAP)
    return body.find("w:sectPr", _NSMAP)


# ---------------------------------------------------------------------------
# Kich ban phau thuat hoan chinh
# ---------------------------------------------------------------------------
def test_dogfood_full_surgical_pipeline(tmp_path):
    docx_path = tmp_path / "bang_diem_dogfood.docx"
    _build_report_card_docx(docx_path)
    allowed_roots = [str(tmp_path)]

    # -----------------------------------------------------------------
    # Buoc 1: inspect - quet tai lieu, trich xuat locators co context_sha256
    # -----------------------------------------------------------------
    report0 = inspect_document(docx_path, work_dir=tmp_path, job_id="job-inspect-0")
    revision_0 = report0["document_revision"]
    assert revision_0.startswith("sha256:")

    toan_locator = _locator(report0, "cell_t1_r2_c2")
    assert toan_locator.kind == "table_cell"
    assert toan_locator.expected_text == "8.5"
    assert toan_locator.context_sha256.startswith("sha256:")

    # -----------------------------------------------------------------
    # Buoc 2: patch-cell - sua diem Toan 8.5 -> 9.0 dung locator ben vung
    # -----------------------------------------------------------------
    patch1 = patch_cell(
        docx_path, toan_locator, "9.0",
        work_dir=tmp_path, allowed_roots=allowed_roots, actor=ACTOR, job_id="job-patch-cell",
    )
    assert patch1.success is True
    assert patch1.outcome == "committed_clean"
    revision_1 = patch1.new_document_revision
    assert revision_1 != revision_0
    assert Document(str(docx_path)).tables[0].cell(1, 1).text == "9.0"

    # -----------------------------------------------------------------
    # Buoc 3: patch-text - sua cum tu trong doan ket luan dung Multi-Run Resolver (khong locator,
    # tim/thay trong toan bo tai lieu)
    # -----------------------------------------------------------------
    patch2 = patch_text(
        docx_path,
        "du dieu kien xet tot nghiep",
        "du dieu kien xet tot nghiep va nhap hoc dung han",
        None, None,
        work_dir=tmp_path, allowed_roots=allowed_roots, actor=ACTOR, job_id="job-patch-text",
    )
    assert patch2.success is True
    assert patch2.outcome == "committed_clean"
    revision_2 = patch2.new_document_revision
    assert revision_2 != revision_1
    assert "nhap hoc dung han" in Document(str(docx_path)).paragraphs[-1].text

    # -----------------------------------------------------------------
    # Buoc 4: set-geometry - chuan hoa kho giay A4, le cong chung (notary) va dan trang
    # (cantSplit/tblHeader) cho bang diem
    # -----------------------------------------------------------------
    geom = apply_geometry(
        docx_path,
        page_size="A4", margins="notary", pagination=True, borders_preset="notary-standard",
        table_locator=None,
        work_dir=tmp_path, allowed_roots=allowed_roots, actor=ACTOR, job_id="job-geometry",
    )
    assert geom.success is True
    assert geom.outcome == "committed_clean"
    revision_3 = geom.new_document_revision
    assert revision_3 != revision_2

    sect_pr = _sect_pr(docx_path)
    pg_sz = sect_pr.find("w:pgSz", _NSMAP)
    assert pg_sz.get(f"{{{WORD_NS}}}w") == "11906"
    assert pg_sz.get(f"{{{WORD_NS}}}h") == "16838"
    pg_mar = sect_pr.find("w:pgMar", _NSMAP)
    assert pg_mar.get(f"{{{WORD_NS}}}left") == "1701"
    assert pg_mar.get(f"{{{WORD_NS}}}top") == "1134"

    root_after_geom = load_document_root(docx_path)
    tbl = root_after_geom.find("w:body", _NSMAP).find("w:tbl", _NSMAP)
    first_row_tr_pr = tbl.find("w:tr", _NSMAP).find("w:trPr", _NSMAP)
    assert first_row_tr_pr.find("w:cantSplit", _NSMAP) is not None
    assert first_row_tr_pr.find("w:tblHeader", _NSMAP) is not None

    # -----------------------------------------------------------------
    # Buoc 5: stamp-ops - chen anh dau/chu ky da duoc whitelist vao cuoi van ban
    # -----------------------------------------------------------------
    asset_path = tmp_path / "assets" / "dau_do.png"
    _make_stamp_png(asset_path)
    register_stamp_asset(asset_path, work_dir=tmp_path, label="Dau do test", registered_by=ACTOR)

    report_before_stamp = inspect_document(docx_path, work_dir=tmp_path, job_id="job-inspect-1")
    assert report_before_stamp["document_revision"] == revision_3
    paragraph_locators = [l for l in report_before_stamp["locators"] if l["kind"] == "paragraph"]
    conclusion_locator = ObjectLocator(**paragraph_locators[-1])  # doan ket luan la doan cuoi cung

    stamp_result = apply_stamp(
        docx_path, asset_path, conclusion_locator, width_mm=30.0, height_mm=20.0,
        work_dir=tmp_path, allowed_roots=allowed_roots, actor=ACTOR, job_id="job-stamp",
    )
    assert stamp_result.success is True
    assert stamp_result.outcome == "committed_clean"
    revision_4 = stamp_result.new_document_revision
    assert revision_4 != revision_3

    with zipfile.ZipFile(docx_path, "r") as zf:
        media_files = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert len(media_files) == 1
    # Van mo lai binh thuong bang python-docx sau khi chen anh (khong lam hong goi OPC).
    assert Document(str(docx_path)).tables[0].cell(1, 1).text == "9.0"

    # -----------------------------------------------------------------
    # Buoc 6: backups - kiem tra danh sach ban sao luu sinh ra sau 4 lan ghi de (patch-cell,
    # patch-text, set-geometry, stamp-ops)
    # -----------------------------------------------------------------
    backups = list_backups(docx_path, work_dir=tmp_path)
    assert len(backups) == 4
    hashes_newest_first = [b["sha256"] for b in backups]
    # Moi ngay truoc: bo sao luu cu nhat (chup truoc buoc patch-cell) phai la sha cua revision_0.
    assert hashes_newest_first[-1] == revision_0.split(":", 1)[1]
    # Bo sao luu moi nhat (chup truoc buoc stamp-ops) phai la sha cua revision_3.
    assert hashes_newest_first[0] == revision_3.split(":", 1)[1]

    # -----------------------------------------------------------------
    # Buoc 7a: restore - khoi phuc ve ngay truoc lan sua gan nhat (truoc stamp-ops)
    # -----------------------------------------------------------------
    latest_backup_id = backups[0]["backup_id"]
    restore_result = restore_backup(
        docx_path, latest_backup_id, work_dir=tmp_path, allowed_roots=allowed_roots, actor=ACTOR,
    )
    assert restore_result["success"] is True
    assert restore_result["outcome"] == "restored_from_backup"
    assert restore_result["restored_from_backup_id"] == latest_backup_id
    assert f"sha256:{restore_result['sha256']}" == revision_3

    with zipfile.ZipFile(docx_path, "r") as zf:
        assert not [n for n in zf.namelist() if n.startswith("word/media/")]  # anh dau da bi go bo
    doc_restored = Document(str(docx_path))
    assert doc_restored.tables[0].cell(1, 1).text == "9.0"  # patch-cell + patch-text van con
    assert "nhap hoc dung han" in doc_restored.paragraphs[-1].text
    assert _sect_pr(docx_path).find("w:pgSz", _NSMAP).get(f"{{{WORD_NS}}}w") == "11906"  # geometry van con

    # -----------------------------------------------------------------
    # Buoc 7b: RESTORE_CONFLICT - mo phong nhan vien sua tay (mo Word/Notepad ngoai broker) roi
    # thu restore lai -> phai bi tu choi, khong duoc am tham ghi de
    # -----------------------------------------------------------------
    hand_edit_marker = "Da duoc nhan vien sua tay ngoai broker, khong qua Word Engine."
    doc_for_hand_edit = Document(str(docx_path))
    doc_for_hand_edit.add_paragraph(hand_edit_marker)
    doc_for_hand_edit.save(str(docx_path))

    backups_after_hand_edit = list_backups(docx_path, work_dir=tmp_path)
    some_backup_id = backups_after_hand_edit[0]["backup_id"]
    with pytest.raises(RestoreConflictError):
        restore_backup(
            docx_path, some_backup_id, work_dir=tmp_path, allowed_roots=allowed_roots, actor=ACTOR,
        )

    # -----------------------------------------------------------------
    # Buoc 8: Kiem chung toan ven - doc lai TOAN BO file .docx hoan chinh bang python-docx, xac
    # nhan 100% noi dung/dinh dang dung nhu thiet ke (tu choi restore o Buoc 7b la TOAN PHAN: noi
    # dung sua tay khong bi mat, khong bi tron lan voi ban backup).
    # -----------------------------------------------------------------
    final_doc = Document(str(docx_path))
    assert final_doc.paragraphs[0].text == "BANG DIEM HOC TAP"
    assert final_doc.tables[0].cell(1, 0).text == "Toan"
    assert final_doc.tables[0].cell(1, 1).text == "9.0"
    assert final_doc.tables[0].cell(2, 1).text == "7.5"
    assert final_doc.tables[0].cell(3, 1).text == "8.0"
    assert "nhap hoc dung han" in final_doc.paragraphs[-2].text
    assert final_doc.paragraphs[-1].text == hand_edit_marker
    final_sect_pr = _sect_pr(docx_path)
    assert final_sect_pr.find("w:pgSz", _NSMAP).get(f"{{{WORD_NS}}}w") == "11906"
    assert final_sect_pr.find("w:pgMar", _NSMAP).get(f"{{{WORD_NS}}}left") == "1701"
