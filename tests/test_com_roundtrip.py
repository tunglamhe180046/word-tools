"""
Tests for core/com_supervisor.py (Phase 4 - Word COM Supervisor & Dual-Layer QA).

Xem docs/plans/DEDICATED_WORD_ENGINE_TOOL_PLAN.md muc 3.5/3.6/4/5.2 va core/com_supervisor.py's
module docstring cho thiet ke. Phan lon test o day KHONG can Word that (mm->points la toan hoc
thuan tuy, Layer 1/Layer 2 chi thao tac tren XML/pixel gia lap trong bo nho, watchdog dung 1 tien
trinh Python thuong lam "tien trinh gia lap WINWORD" thay vi spawn Word that) - dung tinh than
"graceful fallback ... chay duoc ca tren CI hoac may khong co Word" cua nhiem vu. Chi 1 nhom test
(TestRealWordSmoke) can Word that va tu skip neu ComSupervisor.is_available() False hoac
DispatchEx that bai tren may khong cai Word.
"""
from __future__ import annotations

import subprocess
import sys
import time
from pathlib import Path

import psutil
import pytest
from docx import Document
from lxml import etree

from core.com_supervisor import (
    MM_TO_POINTS,
    ComSupervisor,
    VisualDiffViolationError,
    VmlIntegrityError,
    WordComUnavailableError,
    _process_is_same,
    _terminate_owned_process,
    _Watchdog,
    compute_masked_visual_diff,
    inspect_vml_signature,
    is_available,
    load_document_xml_root_from_docx,
    mm_to_points,
    require_masked_visual_diff_clean,
    require_vml_integrity,
    validate_vml_integrity,
)

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---------------------------------------------------------------------------
# 1. mm -> points conversion
# ---------------------------------------------------------------------------
class TestMmToPoints:
    def test_conversion_factor_matches_spec(self):
        assert MM_TO_POINTS == pytest.approx(72.0 / 25.4)

    def test_1mm_converts_to_expected_points(self):
        assert mm_to_points(1.0) == pytest.approx(2.834645669, rel=1e-6)

    def test_30mm_stamp_size_converts_correctly(self):
        assert mm_to_points(30.0) == pytest.approx(30.0 * 72.0 / 25.4, rel=1e-9)

    def test_zero_mm_converts_to_zero_points(self):
        assert mm_to_points(0.0) == 0.0


# ---------------------------------------------------------------------------
# 2. Layer 1: Structural VML Validation
# ---------------------------------------------------------------------------
O_NS = "urn:schemas-microsoft-com:office:office"


def _build_document_root_with_pict(rid: str = "rId4", shape_id: str = "_x0000_i1025") -> etree._Element:
    xml = f"""
    <w:document xmlns:w="{WORD_NS}" xmlns:v="{V_NS}" xmlns:r="{R_NS}" xmlns:o="{O_NS}">
      <w:body>
        <w:p>
          <w:r>
            <w:pict>
              <v:shape id="{shape_id}" type="#_x0000_t75">
                <v:imagedata r:id="{rid}" o:title=""/>
              </v:shape>
            </w:pict>
          </w:r>
        </w:p>
      </w:body>
    </w:document>
    """
    return etree.fromstring(xml.encode("utf-8"))


def _build_document_root_without_pict() -> etree._Element:
    xml = f"""
    <w:document xmlns:w="{WORD_NS}">
      <w:body>
        <w:p><w:r><w:t>No VML here.</w:t></w:r></w:p>
      </w:body>
    </w:document>
    """
    return etree.fromstring(xml.encode("utf-8"))


class TestVmlStructuralValidation:
    def test_intact_vml_round_trip_is_valid(self):
        before = _build_document_root_with_pict()
        after = _build_document_root_with_pict()
        report = validate_vml_integrity(before, after)
        assert report["valid"] is True
        assert report["dropped"] == []
        require_vml_integrity(before, after)  # khong nem loi

    def test_dropped_pict_is_detected(self):
        before = _build_document_root_with_pict()
        after = _build_document_root_without_pict()
        report = validate_vml_integrity(before, after)
        assert report["valid"] is False
        assert any("w:pict" in msg for msg in report["dropped"])
        with pytest.raises(VmlIntegrityError):
            require_vml_integrity(before, after)

    def test_dropped_imagedata_relationship_is_detected(self):
        before = _build_document_root_with_pict(rid="rId4")
        after = _build_document_root_with_pict(rid="rId9")  # cung so pict, khac r:id
        report = validate_vml_integrity(before, after)
        assert report["valid"] is False
        assert any("r:id" in msg for msg in report["dropped"])

    def test_added_vml_does_not_count_as_dropped(self):
        """Them VML moi (khong lam mat cai cu) khong phai vi pham Layer 1 - chi RUNG moi bi bat."""
        before = _build_document_root_without_pict()
        after = _build_document_root_with_pict()
        report = validate_vml_integrity(before, after)
        assert report["valid"] is True

    def test_inspect_vml_signature_reads_single_docx(self, tmp_path):
        docx_path = tmp_path / "with_pict.docx"
        doc = Document()
        doc.add_paragraph("Tai lieu co VML gia lap se duoc ghi de sau.")
        doc.save(docx_path)

        # Ghi de word/document.xml bang ban co VML de kiem tra inspect_vml_signature() doc dung
        # tu chinh file docx tren dia (khong phai tu 1 root da parse san).
        import zipfile

        root = _build_document_root_with_pict()
        new_bytes = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(docx_path, "r") as src:
            items = {item.filename: src.read(item.filename) for item in src.infolist()}
        items["word/document.xml"] = new_bytes
        with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as dst:
            for name, data in items.items():
                dst.writestr(name, data)

        signature = inspect_vml_signature(docx_path)
        assert signature["pict_count"] == 1
        assert signature["shape_count"] == 1
        assert signature["imagedata_count"] == 1

        root_from_disk = load_document_xml_root_from_docx(docx_path)
        assert root_from_disk.find(f"{{{WORD_NS}}}body") is not None


# ---------------------------------------------------------------------------
# 3. Layer 2: Masked Visual Diff
# ---------------------------------------------------------------------------
class _FakePixmap:
    """Doi tuong "Pixmap-like" toi thieu (width/height/n/samples) - du de test
    compute_masked_visual_diff() ma khong can PyMuPDF rasterize tu PDF that."""

    def __init__(self, width: int, height: int, n: int, fill: int = 255):
        self.width = width
        self.height = height
        self.n = n
        self.samples = bytearray([fill]) * (width * height * n)

    def set_pixel(self, x: int, y: int, value: int) -> None:
        offset = (y * self.width + x) * self.n
        for c in range(self.n):
            self.samples[offset + c] = value


class TestMaskedVisualDiff:
    def test_identical_pixmaps_have_zero_diff_ratio(self):
        before = _FakePixmap(10, 10, 3)
        after = _FakePixmap(10, 10, 3)
        report = compute_masked_visual_diff(before, after)
        assert report["pixel_diff_ratio"] == 0.0
        assert report["diff_pixels_count"] == 0
        require_masked_visual_diff_clean(before, after)  # khong nem loi

    def test_diff_inside_authorized_roi_is_excluded(self):
        before = _FakePixmap(10, 10, 3)
        after = _FakePixmap(10, 10, 3)
        after.set_pixel(2, 2, 0)  # thay doi hop le trong vung sua

        report = compute_masked_visual_diff(before, after, authorized_rois=[(0, 0, 5, 5)])
        assert report["pixel_diff_ratio"] == 0.0
        require_masked_visual_diff_clean(before, after, authorized_rois=[(0, 0, 5, 5)])

    def test_diff_outside_authorized_roi_is_detected(self):
        before = _FakePixmap(10, 10, 3)
        after = _FakePixmap(10, 10, 3)
        after.set_pixel(8, 8, 0)  # thay doi NGOAI vung sua

        report = compute_masked_visual_diff(before, after, authorized_rois=[(0, 0, 5, 5)])
        assert report["pixel_diff_ratio"] > 0.0
        assert (8, 8) in report["sample_diff_pixels"]
        with pytest.raises(VisualDiffViolationError):
            require_masked_visual_diff_clean(before, after, authorized_rois=[(0, 0, 5, 5)])

    def test_mismatched_dimensions_raise_value_error(self):
        before = _FakePixmap(10, 10, 3)
        after = _FakePixmap(12, 10, 3)
        with pytest.raises(ValueError):
            compute_masked_visual_diff(before, after)


# ---------------------------------------------------------------------------
# 4. COM Supervisor: watchdog, isolation, graceful unavailability
# ---------------------------------------------------------------------------
def _spawn_dummy_process(sleep_s: float) -> subprocess.Popen:
    return subprocess.Popen([sys.executable, "-c", f"import time; time.sleep({sleep_s})"])


class TestWatchdogAndIsolation:
    def test_watchdog_terminates_owned_process_after_timeout(self):
        proc = _spawn_dummy_process(sleep_s=30)
        try:
            create_time = psutil.Process(proc.pid).create_time()
            watchdog = _Watchdog(proc.pid, create_time, timeout_s=0.3).start()
            time.sleep(1.0)
            assert watchdog.triggered is True
            assert not _process_is_same(proc.pid, create_time)
        finally:
            proc.wait(timeout=10)

    def test_watchdog_does_not_terminate_when_cancelled_before_timeout(self):
        proc = _spawn_dummy_process(sleep_s=2)
        try:
            create_time = psutil.Process(proc.pid).create_time()
            watchdog = _Watchdog(proc.pid, create_time, timeout_s=5.0).start()
            watchdog.cancel()
            assert watchdog.triggered is False
            assert _process_is_same(proc.pid, create_time)
        finally:
            proc.wait(timeout=10)

    def test_terminate_refuses_mismatched_create_time(self):
        """Bat bien an toan cung: khong bao gio kill 1 PID neu create_time khong khop - mo phong
        truong hop PID bi he dieu hanh tai su dung cho 1 tien trinh khac sau khi tien trinh
        supervisor da tu spawn ket thuc truoc do."""
        proc = _spawn_dummy_process(sleep_s=5)
        try:
            real_create_time = psutil.Process(proc.pid).create_time()
            wrong_create_time = real_create_time - 999.0
            result = _terminate_owned_process(proc.pid, wrong_create_time)
            assert result is False
            assert proc.poll() is None  # van con song - khong bi dung nham
        finally:
            proc.terminate()
            proc.wait(timeout=10)

    def test_terminate_returns_false_for_nonexistent_pid(self):
        # PID rat lon it kha nang trung mot tien trinh dang chay that tren may test.
        assert _terminate_owned_process(999999, time.time()) is False

    def test_process_is_same_false_for_nonexistent_pid(self):
        assert _process_is_same(999999, time.time()) is False


class TestGracefulUnavailability:
    def test_is_available_false_when_not_windows(self, monkeypatch):
        monkeypatch.setattr("core.com_supervisor.platform.system", lambda: "Linux")
        assert is_available() is False

    def test_is_available_false_when_pywin32_missing(self, monkeypatch):
        monkeypatch.setattr("core.com_supervisor.is_windows", lambda: True)
        import builtins

        real_import = builtins.__import__

        def _fake_import(name, *args, **kwargs):
            if name in ("win32com.client", "pythoncom"):
                raise ImportError(f"simulated missing module: {name}")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", _fake_import)
        assert is_available() is False

    def test_export_pdf_raises_word_com_unavailable_when_not_available(self, tmp_path, monkeypatch):
        monkeypatch.setattr("core.com_supervisor.is_available", lambda: False)
        docx_path = tmp_path / "sample.docx"
        Document().save(docx_path)

        supervisor = ComSupervisor()
        with pytest.raises(WordComUnavailableError):
            supervisor.export_pdf(docx_path, tmp_path / "sample.pdf")

    def test_export_pdf_missing_docx_raises_file_not_found(self, tmp_path):
        supervisor = ComSupervisor()
        with pytest.raises(FileNotFoundError):
            supervisor.export_pdf(tmp_path / "missing.docx", tmp_path / "missing.pdf")

    def test_run_job_never_reaches_fn_when_unavailable(self, monkeypatch):
        monkeypatch.setattr("core.com_supervisor.is_available", lambda: False)
        supervisor = ComSupervisor()
        called = {"value": False}

        def _fn(app):
            called["value"] = True
            return app

        with pytest.raises(WordComUnavailableError):
            supervisor.run_job(_fn)
        assert called["value"] is False


# ---------------------------------------------------------------------------
# 5. Optional real-Word smoke test - skipped automatically when Word COM is not
#    actually usable on this machine (CI / machine without Word installed).
# ---------------------------------------------------------------------------
class TestRealWordSmoke:
    @pytest.mark.skipif(not is_available(), reason="Word COM khong kha dung tren may nay.")
    def test_export_pdf_produces_real_pdf(self, tmp_path):
        docx_path = tmp_path / "smoke.docx"
        doc = Document()
        doc.add_paragraph("Word COM Supervisor smoke test.")
        doc.save(docx_path)

        supervisor = ComSupervisor(watchdog_timeout_s=60.0)
        try:
            pdf_path = supervisor.export_pdf(docx_path, tmp_path / "smoke.pdf")
        except WordComUnavailableError:
            pytest.skip("Word chua duoc cai dat that su tren may nay (DispatchEx that bai).")
        assert pdf_path.exists()
        assert pdf_path.stat().st_size > 0

        import pymupdf

        pdf_doc = pymupdf.open(str(pdf_path))
        try:
            assert pdf_doc.page_count >= 1
        finally:
            pdf_doc.close()
