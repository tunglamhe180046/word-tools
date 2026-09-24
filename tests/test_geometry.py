"""
Tests for core/geometry_styler.py (Phase 2 - Geometry & Layout Standardization) and cli.py's
`set-geometry` subcommand + the --json argparse-error contract added alongside it.

See docs/plans/DEDICATED_WORD_ENGINE_TOOL_PLAN.md muc 4 ("Phase 2") and
core/geometry_styler.py's module docstring for the design. No mocking, real files on disk under
tmp_path, real SHA-256/XML comparisons - same style as tests/test_surgical_safety.py and
tests/test_cli.py. Every test passes an explicit work_dir=tmp_path so `.jarvis/` never touches the
real repo's `.jarvis/`.
"""
from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document

from core.geometry_styler import apply_geometry
from core.inspector import WORD_NS, DocumentDriftError, ObjectLocator, inspect_document, load_document_root

_NSMAP = {"w": WORD_NS}
CLI_PATH = Path(__file__).resolve().parent.parent / "cli.py"

COMMON_KWARGS = dict(actor="tester", job_id="job1")


def _w(tag: str) -> str:
    return f"{{{WORD_NS}}}{tag}"


# ---------------------------------------------------------------------------
# Fixtures & small helpers
# ---------------------------------------------------------------------------
def _build_table_docx(path: Path, rows: int = 3, cols: int = 2, with_heading: bool = True) -> None:
    doc = Document()
    if with_heading:
        doc.add_paragraph("Tieu de bang so lieu.")
    table = doc.add_table(rows=rows, cols=cols)
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = f"O {r}-{c}"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _build_two_table_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Doan mo dau.")
    t1 = doc.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "A1"
    t1.cell(0, 1).text = "A2"
    t1.cell(1, 0).text = "A3"
    t1.cell(1, 1).text = "A4"
    doc.add_paragraph("Doan giua.")
    t2 = doc.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "B1"
    t2.cell(0, 1).text = "B2"
    t2.cell(1, 0).text = "B3"
    t2.cell(1, 1).text = "B4"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _part_hashes(docx_path: Path) -> dict:
    with zipfile.ZipFile(docx_path, "r") as zf:
        return {name: sha256(zf.read(name)).hexdigest() for name in zf.namelist()}


def _run_cli(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(CLI_PATH), *args], capture_output=True, text=True)


# ---------------------------------------------------------------------------
# Nhom 3: Page Setup
# ---------------------------------------------------------------------------
def test_apply_geometry_sets_page_size_a4(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    result = apply_geometry(
        docx_path, page_size="A4", margins=None, pagination=False, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )
    assert result.success is True

    body = load_document_root(docx_path).find("w:body", _NSMAP)
    pg_sz = body.find("w:sectPr", _NSMAP).find("w:pgSz", _NSMAP)
    assert pg_sz.get(_w("w")) == "11906"
    assert pg_sz.get(_w("h")) == "16838"


def test_apply_geometry_sets_page_size_a5(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    apply_geometry(
        docx_path, page_size="A5", margins=None, pagination=False, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    body = load_document_root(docx_path).find("w:body", _NSMAP)
    pg_sz = body.find("w:sectPr", _NSMAP).find("w:pgSz", _NSMAP)
    assert pg_sz.get(_w("w")) == "8391"
    assert pg_sz.get(_w("h")) == "11906"


def test_apply_geometry_sets_notary_margins(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    apply_geometry(
        docx_path, page_size=None, margins="notary", pagination=False, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    body = load_document_root(docx_path).find("w:body", _NSMAP)
    pg_mar = body.find("w:sectPr", _NSMAP).find("w:pgMar", _NSMAP)
    assert pg_mar.get(_w("left")) == "1701"
    assert pg_mar.get(_w("right")) == "850"
    assert pg_mar.get(_w("top")) == "1134"
    assert pg_mar.get(_w("bottom")) == "1134"


def test_apply_geometry_sets_compact_margins(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    apply_geometry(
        docx_path, page_size=None, margins="compact", pagination=False, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    body = load_document_root(docx_path).find("w:body", _NSMAP)
    pg_mar = body.find("w:sectPr", _NSMAP).find("w:pgMar", _NSMAP)
    for edge in ("left", "right", "top", "bottom"):
        assert pg_mar.get(_w(edge)) == "720"


# ---------------------------------------------------------------------------
# Nhom 1: Pagination
# ---------------------------------------------------------------------------
def test_apply_geometry_pagination_sets_cantsplit_all_rows_and_tblheader_first_row_only(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=3, cols=2, with_heading=True)

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=True, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    body = load_document_root(docx_path).find("w:body", _NSMAP)
    tbl = body.findall("w:tbl", _NSMAP)[0]
    rows = tbl.findall("w:tr", _NSMAP)
    assert len(rows) == 3
    for i, tr in enumerate(rows):
        tr_pr = tr.find("w:trPr", _NSMAP)
        assert tr_pr is not None
        assert tr_pr.find("w:cantSplit", _NSMAP) is not None
        has_header = tr_pr.find("w:tblHeader", _NSMAP) is not None
        assert has_header == (i == 0), f"row {i}: tblHeader present={has_header}"

    # Doan van ngay truoc bang phai co w:keepNext (xem module docstring ve ten the that su).
    heading_p_pr = body.findall("w:p", _NSMAP)[0].find("w:pPr", _NSMAP)
    assert heading_p_pr is not None
    assert heading_p_pr.find("w:keepNext", _NSMAP) is not None


def test_apply_geometry_pagination_without_preceding_paragraph_does_not_fail(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=2, cols=2, with_heading=False)

    result = apply_geometry(
        docx_path, page_size=None, margins=None, pagination=True, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )
    assert result.success is True


# ---------------------------------------------------------------------------
# Nhom 2: Borders
# ---------------------------------------------------------------------------
def test_apply_geometry_notary_standard_borders(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=2, cols=2)

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=False, borders_preset="notary-standard",
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    tbl = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)[0]
    tbl_borders = tbl.find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP)
    assert tbl_borders.find("w:top", _NSMAP).get(_w("val")) == "single"
    assert tbl_borders.find("w:top", _NSMAP).get(_w("sz")) == "4"
    assert tbl_borders.find("w:insideH", _NSMAP).get(_w("val")) == "single"
    assert tbl_borders.find("w:insideH", _NSMAP).get(_w("sz")) == "2"
    assert tbl_borders.find("w:insideV", _NSMAP).get(_w("val")) == "nil"

    # Moi o truc tiep deu co w:tcBorders rieng, cung preset (de bai nhiem vu: "va w:tcBorders
    # (tung o)").
    for tr in tbl.findall("w:tr", _NSMAP):
        for tc in tr.findall("w:tc", _NSMAP):
            tc_borders = tc.find("w:tcPr", _NSMAP).find("w:tcBorders", _NSMAP)
            assert tc_borders is not None
            assert tc_borders.find("w:insideV", _NSMAP).get(_w("val")) == "nil"
            assert tc_borders.find("w:top", _NSMAP).get(_w("val")) == "single"


def test_apply_geometry_borders_none_sets_nil_on_every_edge(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=2, cols=2)

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=False, borders_preset="none",
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    tbl = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)[0]
    tbl_borders = tbl.find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        assert tbl_borders.find(f"w:{edge}", _NSMAP).get(_w("val")) == "nil"


def test_apply_geometry_borders_horizontal_only(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=2, cols=2)

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=False, borders_preset="horizontal-only",
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    tbl_borders = (
        load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)[0]
        .find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP)
    )
    assert tbl_borders.find("w:insideH", _NSMAP).get(_w("val")) == "single"
    assert tbl_borders.find("w:left", _NSMAP).get(_w("val")) == "nil"
    assert tbl_borders.find("w:right", _NSMAP).get(_w("val")) == "nil"
    assert tbl_borders.find("w:insideV", _NSMAP).get(_w("val")) == "nil"


def test_apply_geometry_invalid_borders_preset_raises_value_error(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    with pytest.raises(ValueError):
        apply_geometry(
            docx_path, page_size=None, margins=None, pagination=False, borders_preset="rainbow",
            table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
        )


# ---------------------------------------------------------------------------
# 3 nhom hoan toan doc lap: khong lam mat cau hinh cua nhau, du goi theo thu tu nao
# ---------------------------------------------------------------------------
def test_apply_geometry_borders_and_pagination_do_not_clobber_each_other(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=2, cols=2, with_heading=False)

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=True, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)],
        actor="tester", job_id="job-pagination",
    )
    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=False, borders_preset="all",
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)],
        actor="tester", job_id="job-borders",
    )

    tbl = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)[0]
    # Pagination tu lan patch truoc van con nguyen sau khi set borders o lan patch sau.
    for tr in tbl.findall("w:tr", _NSMAP):
        assert tr.find("w:trPr", _NSMAP).find("w:cantSplit", _NSMAP) is not None
    assert tbl.find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP) is not None

    # Nguoc lai: set pagination sau khi da co borders khong lam mat borders.
    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=True, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)],
        actor="tester", job_id="job-pagination-2",
    )
    tbl2 = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)[0]
    assert tbl2.find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP) is not None
    for tr in tbl2.findall("w:tr", _NSMAP):
        assert tr.find("w:trPr", _NSMAP).find("w:cantSplit", _NSMAP) is not None


def test_apply_geometry_leaves_untouched_opc_parts_raw_sha256_identical(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)
    before = _part_hashes(docx_path)

    apply_geometry(
        docx_path, page_size="A4", margins="notary", pagination=True, borders_preset="all",
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    after = _part_hashes(docx_path)
    assert set(before) == set(after)
    for name in before:
        if name == "word/document.xml":
            continue
        assert after[name] == before[name], f"part {name} changed unexpectedly"
    assert after["word/document.xml"] != before["word/document.xml"]


# ---------------------------------------------------------------------------
# table_locator: chi ap dung cho DUNG 1 bang khi tai lieu co nhieu hon 1 bang
# ---------------------------------------------------------------------------
def test_apply_geometry_with_table_locator_scopes_to_target_table_only(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_two_table_docx(docx_path)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = ObjectLocator(**next(loc for loc in report["locators"] if loc["object_id"] == "cell_t2_r1_c1"))

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=True, borders_preset="all",
        table_locator=locator, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    tables = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)
    assert len(tables) == 2

    t1 = tables[0]
    t1_pr = t1.find("w:tblPr", _NSMAP)
    assert t1_pr is None or t1_pr.find("w:tblBorders", _NSMAP) is None
    for tr in t1.findall("w:tr", _NSMAP):
        assert tr.find("w:trPr", _NSMAP) is None

    t2 = tables[1]
    assert t2.find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP) is not None
    for tr in t2.findall("w:tr", _NSMAP):
        assert tr.find("w:trPr", _NSMAP).find("w:cantSplit", _NSMAP) is not None


def test_apply_geometry_without_table_locator_applies_to_every_table(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_two_table_docx(docx_path)

    apply_geometry(
        docx_path, page_size=None, margins=None, pagination=True, borders_preset=None,
        table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
    )

    tables = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)
    assert len(tables) == 2
    for tbl in tables:
        for tr in tbl.findall("w:tr", _NSMAP):
            assert tr.find("w:trPr", _NSMAP).find("w:cantSplit", _NSMAP) is not None


def test_apply_geometry_with_stale_table_locator_raises_document_drift(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_two_table_docx(docx_path)
    report = inspect_document(docx_path, work_dir=tmp_path, job_id="job1")
    locator = ObjectLocator(**next(loc for loc in report["locators"] if loc["object_id"] == "cell_t1_r1_c1"))

    doc = Document(str(docx_path))
    doc.add_paragraph("Nguoi dung sua tay them doan nay.")
    doc.save(docx_path)
    before_bytes = docx_path.read_bytes()

    with pytest.raises(DocumentDriftError):
        apply_geometry(
            docx_path, page_size=None, margins=None, pagination=True, borders_preset=None,
            table_locator=locator, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
        )
    assert docx_path.read_bytes() == before_bytes


def test_apply_geometry_requires_at_least_one_option(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    with pytest.raises(ValueError):
        apply_geometry(
            docx_path, page_size=None, margins=None, pagination=False, borders_preset=None,
            table_locator=None, work_dir=tmp_path, allowed_roots=[str(tmp_path)], **COMMON_KWARGS,
        )


# ---------------------------------------------------------------------------
# CLI: set-geometry --json contract
# ---------------------------------------------------------------------------
def test_cli_set_geometry_json_contract(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, rows=2, cols=2)

    proc = _run_cli(
        "set-geometry", str(docx_path),
        "--page-size", "A4", "--margins", "notary", "--pagination", "--borders", "notary-standard",
        "--json", "--work-dir", str(tmp_path),
    )

    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.count("\n") == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["outcome"] == "committed_clean"

    body = load_document_root(docx_path).find("w:body", _NSMAP)
    assert body.find("w:sectPr", _NSMAP).find("w:pgSz", _NSMAP).get(_w("w")) == "11906"
    tbl = body.findall("w:tbl", _NSMAP)[0]
    assert tbl.find("w:tr", _NSMAP).find("w:trPr", _NSMAP).find("w:cantSplit", _NSMAP) is not None
    assert tbl.find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP) is not None


def test_cli_set_geometry_with_table_id_scopes_target_table(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_two_table_docx(docx_path)

    proc = _run_cli(
        "set-geometry", str(docx_path),
        "--borders", "all", "--table-id", "cell_t2_r1_c1",
        "--json", "--work-dir", str(tmp_path),
    )

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["success"] is True

    tables = load_document_root(docx_path).find("w:body", _NSMAP).findall("w:tbl", _NSMAP)
    t1_pr = tables[0].find("w:tblPr", _NSMAP)
    assert t1_pr is None or t1_pr.find("w:tblBorders", _NSMAP) is None
    assert tables[1].find("w:tblPr", _NSMAP).find("w:tblBorders", _NSMAP) is not None


def test_cli_set_geometry_unknown_table_id_fails_cleanly(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    proc = _run_cli(
        "set-geometry", str(docx_path), "--pagination", "--table-id", "does-not-exist",
        "--json", "--work-dir", str(tmp_path),
    )

    assert proc.returncode == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "LocatorNotFoundError"


# ---------------------------------------------------------------------------
# CLI: --json argparse-error contract (Phan 1 cua nhiem vu)
# ---------------------------------------------------------------------------
def test_cli_json_argparse_error_missing_required_arguments(tmp_path):
    proc = _run_cli("patch-cell", "--json")

    assert proc.returncode == 1
    assert proc.stdout.count("\n") == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "ArgumentError"
    assert payload["reason"]
    assert proc.stderr == ""


def test_cli_json_argparse_error_unknown_subcommand(tmp_path):
    proc = _run_cli("not-a-real-command", "--json")

    assert proc.returncode == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "ArgumentError"


def test_cli_json_argparse_error_invalid_choice_value(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path)

    proc = _run_cli(
        "set-geometry", str(docx_path), "--borders", "not-a-real-preset",
        "--json", "--work-dir", str(tmp_path),
    )

    assert proc.returncode == 1
    payload = json.loads(proc.stdout)
    assert payload["success"] is False
    assert payload["error"] == "ArgumentError"


def test_cli_argparse_error_without_json_keeps_default_argparse_behavior(tmp_path):
    proc = _run_cli("patch-cell")

    assert proc.returncode == 2
    assert proc.stdout == ""
    assert proc.stderr != ""


def test_apply_geometry_paragraph_alignment(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _build_table_docx(docx_path, with_heading=True)

    report = inspect_document(docx_path, work_dir=tmp_path, job_id="inspect-align")
    para_loc = next(loc for loc in report["locators"] if loc["kind"] == "paragraph")

    proc = _run_cli(
        "set-geometry", str(docx_path), "--align", "center",
        "--target-id", para_loc["object_id"], "--json", "--work-dir", str(tmp_path),
    )
    assert proc.returncode == 0
    payload = json.loads(proc.stdout)
    assert payload["success"] is True
    assert payload["outcome"] == "committed_clean"

    root = load_document_root(docx_path)
    body = root.find("w:body", _NSMAP)
    p = body.findall("w:p", _NSMAP)[0]
    p_pr = p.find("w:pPr", _NSMAP)
    assert p_pr is not None
    jc = p_pr.find("w:jc", _NSMAP)
    assert jc is not None
    assert jc.attrib[f"{{{WORD_NS}}}val"] == "center"

